#!/usr/bin/env python3
"""Mars 1-lb catalog matrix. Index n = 1+0 (never 0)."""

from openpyxl import Workbook
from openpyxl.styles import Font, Fill, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

wb = Workbook()

# --- styles ---
title_font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
label_font = Font(name="Arial", size=10, bold=True)
body_font = Font(name="Arial", size=10)
blue_input = Font(name="Arial", size=10, color="0000FF")
black_calc = Font(name="Arial", size=10, color="000000")
note_font = Font(name="Arial", size=9, italic=True, color="333333")

fill_header = PatternFill("solid", fgColor="1B3A4B")
fill_title = PatternFill("solid", fgColor="0D1B2A")
fill_n = PatternFill("solid", fgColor="E8F1F5")
fill_yellow = PatternFill("solid", fgColor="FFF3B0")
fill_alt = PatternFill("solid", fgColor="F4F7F8")
fill_core = PatternFill("solid", fgColor="D8E8D8")
fill_enab = PatternFill("solid", fgColor="D6E4F0")
fill_opt = PatternFill("solid", fgColor="F5E6D3")
fill_algo = PatternFill("solid", fgColor="1B3A4B")

thin = Border(
    left=Side(style="thin", color="B0BEC5"),
    right=Side(style="thin", color="B0BEC5"),
    top=Side(style="thin", color="B0BEC5"),
    bottom=Side(style="thin", color="B0BEC5"),
)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center", wrap_text=True)

# ============================================================
# Sheet 1 — Algorithm
# ============================================================
ws0 = wb.active
ws0.title = "Algorithm"

ws0.merge_cells("A1:F1")
ws0["A1"] = "MARS 1-LB CATALOG  ·  n = 1 + 0"
ws0["A1"].font = title_font
ws0["A1"].fill = fill_title
ws0["A1"].alignment = Alignment(horizontal="left", vertical="center")
ws0.row_dimensions[1].height = 28

ws0.merge_cells("A2:F2")
ws0["A2"] = "Catalog index never starts at zero. First slot is n = 1. Offset is 0. n = 1 + 0."
ws0["A2"].font = Font(name="Arial", size=11, italic=True, color="1B3A4B")
ws0.row_dimensions[2].height = 20

ws0["A4"] = "RULE"
ws0["A4"].font = header_font
ws0["A4"].fill = fill_algo
ws0.merge_cells("B4:F4")
ws0["B4"] = "n ≠ 0.  n = 1 + 0.  Every catalog row is addressed by a positive integer index."
ws0["B4"].font = Font(name="Arial", size=11, bold=True)
ws0["B4"].fill = fill_yellow

rows_algo = [
    ("Symbol", "Meaning", "Value / formula", "Notes", "", ""),
    ("n", "Catalog index", "1, 2, 3, …", "Row address. Never 0.", "", ""),
    ("n0", "Offset", "0", "Added so the first item is 1, not 0.", "", ""),
    ("n_start", "First legal index", "=1+0", "Hard rule of this matrix.", "", ""),
    ("m_i", "Catalog pounds of item i", "input (blue)", "Mass expressed in 1-lb units. May be decimal.", "", ""),
    ("p_i", "Packet ceiling", "=CEILING(m_i,1)", "If you later packetize. Catalog itself allows fractions.", "", ""),
    ("k_i", "Assembly width", "input", "How many catalog-lb units form one useful assembled object.", "", ""),
    ("class_i", "Could-go class", "Core / Enabling / Optional", "Not a flight commitment.", "", ""),
]

for r, row in enumerate(rows_algo, start=6):
    for c, val in enumerate(row, start=1):
        cell = ws0.cell(r, c, val)
        cell.font = header_font if r == 6 else body_font
        cell.fill = fill_header if r == 6 else PatternFill()
        cell.alignment = left
        cell.border = thin

ws0["C9"] = "=1+0"
ws0["C9"].font = black_calc
ws0["C9"].fill = fill_yellow

ws0.merge_cells("A16:F16")
ws0["A16"] = "INDEX CONSTRUCTION"
ws0["A16"].font = header_font
ws0["A16"].fill = fill_header

ws0.merge_cells("A17:F20")
ws0["A17"] = (
    "Row i on the Catalog sheet is addressed as n = i, with i starting at 1.\n"
    "There is no row 0 and no item 0. Empty is not an index.\n"
    "n = 1 + 0 is the generator: start at one, add a zero offset, never decrement to zero.\n"
    "New items append at n_max + 1. Deleting an item retires the number; it is not reused as 0."
)
ws0["A17"].alignment = Alignment(wrap_text=True, vertical="top")
ws0["A17"].font = body_font
ws0.row_dimensions[17].height = 20
ws0.row_dimensions[18].height = 20
ws0.row_dimensions[19].height = 20
ws0.row_dimensions[20].height = 20

ws0.merge_cells("A22:F22")
ws0["A22"] = "This workbook catalogs infrastructure that COULD go to Mars. It is not a flight manifest."
ws0["A22"].font = note_font

for col, w in enumerate([14, 28, 28, 42, 14, 14], start=1):
    ws0.column_dimensions[get_column_letter(col)].width = w

# ============================================================
# Sheet 2 — Catalog matrix
# ============================================================
ws = wb.create_sheet("Catalog")

ws.merge_cells("A1:L1")
ws["A1"] = "MARS INFRASTRUCTURE CATALOG  ·  unit = 1 pound  ·  index n = 1 + 0"
ws["A1"].font = title_font
ws["A1"].fill = fill_title
ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
ws.row_dimensions[1].height = 26

ws.merge_cells("A2:L2")
ws["A2"] = "Could-go inventory only. Blue cells are inputs. Black cells are formulas. n never equals 0."
ws["A2"].font = Font(name="Arial", size=10, italic=True, color="1B3A4B")

headers = [
    "n",
    "Category",
    "Item",
    "Catalog lb  m_i",
    "Packet ceil  p_i",
    "Assembly k",
    "Could-go class",
    "Depends on (n)",
    "Enables",
    "Notes",
    "Running lb",
    "Running packets",
]

for c, h in enumerate(headers, start=1):
    cell = ws.cell(3, c, h)
    cell.font = header_font
    cell.fill = fill_header
    cell.alignment = center
    cell.border = thin
ws.row_dimensions[3].height = 32
ws.auto_filter.ref = "A3:L53"
ws.freeze_panes = "A4"

# Seed catalog — n starts at 1
items = [
    (1, "Habitat", "Unpressurized deck tile", 1.0, 4, "Core", "—", "Floor / wall grid", "1-lb base plate"),
    (2, "Habitat", "Pressure panel", 8.0, 4, "Core", "1", "Livable volume", "Needs seals + frame"),
    (3, "Habitat", "Airlock collar", 18.0, 1, "Enabling", "2", "EVA / transfer", ""),
    (4, "Habitat", "Window / viewport brick", 6.0, 1, "Optional", "2", "Sightline", ""),
    (5, "Power", "Solar brick", 1.0, 8, "Core", "—", "Trickle charge", "Dust-tolerant face"),
    (6, "Power", "Battery brick", 2.0, 8, "Core", "5", "Night ops", ""),
    (7, "Power", "Power bus node", 3.0, 2, "Enabling", "5,6", "Distribution", ""),
    (8, "Power", "Small RTG / nuke brick (concept)", 25.0, 1, "Optional", "7", "Base load", "Catalog only — not a build plan"),
    (9, "Mobility", "Wheel / drive packet", 4.0, 4, "Enabling", "10", "Motion", ""),
    (10, "Mobility", "Small rover chassis", 25.0, 1, "Enabling", "9,11", "Local haul", ""),
    (11, "Mobility", "Compute / nav brick", 3.0, 1, "Enabling", "—", "Drive + pose", ""),
    (12, "Mobility", "Manipulator arm segment", 7.0, 3, "Optional", "10", "Pick / place", ""),
    (13, "Life support", "Water recycler brick", 12.0, 1, "Core", "6,7", "Crew days", ""),
    (14, "Life support", "O2 / CO2 cartridge", 2.0, 6, "Core", "13", "Breathing loop", ""),
    (15, "Life support", "Tank / bladder", 5.0, 2, "Enabling", "13", "Store water", ""),
    (16, "Life support", "Thermal garment pack", 4.0, 2, "Optional", "—", "EVA / night", ""),
    (17, "Comms", "Surface relay", 15.0, 1, "Core", "6,7", "Link home / site", ""),
    (18, "Comms", "Local mesh node", 1.0, 8, "Enabling", "17", "Packet radio grid", "1-lb radio"),
    (19, "Comms", "Nav beacon", 2.0, 4, "Enabling", "18", "Site fix", ""),
    (20, "Manufacture", "ISRU feedstock bag", 1.0, 20, "Enabling", "—", "Print / sinter", "Regolith or resin"),
    (21, "Manufacture", "Print head brick", 9.0, 1, "Optional", "6,20", "Make parts", ""),
    (22, "Manufacture", "Sinter / kiln brick", 14.0, 1, "Optional", "6,20", "Brick / tile from dirt", ""),
    (23, "Science", "Weather / dust station", 3.0, 1, "Enabling", "6,18", "Site data", ""),
    (24, "Science", "Seismic pin", 1.0, 4, "Optional", "18", "Ground motion", ""),
    (25, "Science", "Sample cache tube", 1.0, 12, "Optional", "12", "Return / store", ""),
    (26, "Safety", "Repair tool kit", 6.0, 1, "Core", "—", "Keep other items alive", ""),
    (27, "Safety", "Patch / seal kit", 2.0, 4, "Core", "2", "Pressure repair", ""),
    (28, "Safety", "Fire / dust suppress pack", 5.0, 2, "Enabling", "—", "Site safety", ""),
    (29, "Safety", "Med / trauma brick", 8.0, 1, "Optional", "—", "Crew care", ""),
    (30, "Habitat", "Sleep / berth frame", 16.0, 1, "Optional", "1,2", "Crew rest", ""),
    (31, "Habitat", "Galley brick", 11.0, 1, "Optional", "6,15", "Food prep", ""),
    (32, "Power", "Cable spool 50 m", 4.0, 2, "Enabling", "7", "Run power", ""),
    (33, "Mobility", "Trailer / pallet", 10.0, 2, "Optional", "10", "Haul packets", ""),
    (34, "Manufacture", "Fastener / interface pack", 1.0, 10, "Core", "—", "Join any two items", "Standard 1-lb join kit"),
    (35, "Comms", "High-gain dish segment", 12.0, 1, "Optional", "17", "Earth link margin", ""),
    (36, "Science", "Camera / mast brick", 4.0, 1, "Optional", "11", "Survey", ""),
    (37, "Life support", "Plant / grow tray", 6.0, 2, "Optional", "5,15", "Fresh mass", ""),
    (38, "Habitat", "Radiation cover tile", 3.0, 8, "Enabling", "1", "Dose cut", ""),
    (39, "Safety", "Beacon / find-me", 1.0, 4, "Enabling", "18", "Lost-crew / lost-packet", ""),
    (40, "Manufacture", "Spare interface plate", 1.0, 16, "Core", "34", "Grow the grid", ""),
    (41, "Power", "Dust-wiper module", 2.0, 2, "Optional", "5", "Keep solar alive", ""),
    (42, "Mobility", "Leg / walk packet (bot)", 8.0, 4, "Optional", "11", "Legged scout", "Fits Tesla-Bot-class thinking"),
    (43, "Science", "G-force / IMU brick", 1.0, 2, "Optional", "11", "Motion log", "Nanocraft-adjacent"),
    (44, "Habitat", "Vestibule tunnel ring", 20.0, 1, "Optional", "3", "Connect volumes", ""),
    (45, "Life support", "Waste lock brick", 9.0, 1, "Enabling", "13", "Close the loop", ""),
    (46, "Comms", "Store-and-forward brick", 2.0, 2, "Optional", "18", "Delay-tolerant net", ""),
    (47, "Safety", "Anchor / guy kit", 3.0, 4, "Enabling", "1", "Wind / storm hold-down", ""),
    (48, "Manufacture", "Cutting / drill brick", 7.0, 1, "Optional", "6", "Modify site / parts", ""),
    (49, "Power", "Dump load / heater brick", 3.0, 2, "Optional", "7", "Use surplus / keep warm", ""),
    (50, "Habitat", "Work surface / bench", 8.0, 1, "Optional", "1,34", "Assemble other n", ""),
]

for i, it in enumerate(items):
    r = 4 + i
    n, cat, name, m, k, klass, depends, enables, notes = it
    ws.cell(r, 1, n).font = black_calc
    ws.cell(r, 1).fill = fill_n
    ws.cell(r, 2, cat).font = body_font
    ws.cell(r, 3, name).font = body_font
    c4 = ws.cell(r, 4, m)
    c4.font = blue_input
    c4.number_format = "0.00"
    c5 = ws.cell(r, 5, f"=CEILING(D{r},1)")
    c5.font = black_calc
    c5.number_format = "0"
    c6 = ws.cell(r, 6, k)
    c6.font = blue_input
    ws.cell(r, 7, klass).font = body_font
    ws.cell(r, 8, depends).font = blue_input
    ws.cell(r, 9, enables).font = body_font
    ws.cell(r, 10, notes).font = note_font
    if i == 0:
        ws.cell(r, 11, f"=D{r}").font = black_calc
        ws.cell(r, 12, f"=E{r}").font = black_calc
    else:
        ws.cell(r, 11, f"=K{r-1}+D{r}").font = black_calc
        ws.cell(r, 12, f"=L{r-1}+E{r}").font = black_calc
    ws.cell(r, 11).number_format = "#,##0.00"
    ws.cell(r, 12).number_format = "#,##0"
    fill = fill_alt if i % 2 else PatternFill()
    if klass == "Core":
        ws.cell(r, 7).fill = fill_core
    elif klass == "Enabling":
        ws.cell(r, 7).fill = fill_enab
    else:
        ws.cell(r, 7).fill = fill_opt
    for c in range(1, 13):
        ws.cell(r, c).border = thin
        ws.cell(r, c).alignment = left if c in (2, 3, 8, 9, 10) else center
    ws.row_dimensions[r].height = 18

# blank rows 54-63 ready for n=51+
for i in range(50, 60):
    r = 4 + i
    n = i + 1
    ws.cell(r, 1, n).font = black_calc
    ws.cell(r, 1).fill = fill_n
    ws.cell(r, 4).font = blue_input
    ws.cell(r, 4).number_format = "0.00"
    ws.cell(r, 5, f'=IF(D{r}="","",CEILING(D{r},1))').font = black_calc
    ws.cell(r, 11, f'=IF(D{r}="","",K{r-1}+D{r})').font = black_calc
    ws.cell(r, 12, f'=IF(E{r}="","",L{r-1}+E{r})').font = black_calc
    ws.cell(r, 11).number_format = "#,##0.00"
    ws.cell(r, 12).number_format = "#,##0"
    for c in range(1, 13):
        ws.cell(r, c).border = thin
        ws.cell(r, c).alignment = center
    ws.row_dimensions[r].height = 18

# totals
tr = 65
ws.cell(tr, 1, "").fill = fill_header
ws.merge_cells("A65:C65")
ws["A65"] = "TOTALS (filled rows only)"
ws["A65"].font = header_font
ws["A65"].fill = fill_header
ws["A65"].alignment = left
ws["D65"] = '=SUMIF(D4:D63,">0")'
ws["D65"].font = black_calc
ws["D65"].number_format = "#,##0.00"
ws["D65"].fill = fill_yellow
ws["E65"] = '=SUMIF(E4:E63,">0")'
ws["E65"].font = black_calc
ws["E65"].number_format = "#,##0"
ws["E65"].fill = fill_yellow
ws["K65"] = "=K53"
ws["K65"].font = black_calc
ws["K65"].number_format = "#,##0.00"
ws["L65"] = "=L53"
ws["L65"].font = black_calc
ws["L65"].number_format = "#,##0"
for c in range(1, 13):
    ws.cell(tr, c).fill = fill_header if c <= 3 else fill_yellow
    ws.cell(tr, c).border = thin
    ws.cell(tr, c).font = header_font if c <= 3 else black_calc

ws["A67"] = "n_start"
ws["B67"] = "=1+0"
ws["B67"].font = black_calc
ws["B67"].fill = fill_yellow
ws["C67"] = "must equal 1 — generator n = 1 + 0"
ws["C67"].font = note_font
ws["A68"] = "n_max filled"
ws["B68"] = "=COUNTA(C4:C53)"
ws["B68"].font = black_calc
ws["A69"] = "Core catalog lb"
ws["B69"] = '=SUMIF(G4:G53,"Core",D4:D53)'
ws["B69"].font = black_calc
ws["B69"].number_format = "#,##0.00"
ws["A70"] = "Enabling catalog lb"
ws["B70"] = '=SUMIF(G4:G53,"Enabling",D4:D53)'
ws["B70"].font = black_calc
ws["B70"].number_format = "#,##0.00"
ws["A71"] = "Optional catalog lb"
ws["B71"] = '=SUMIF(G4:G53,"Optional",D4:D53)'
ws["B71"].font = black_calc
ws["B71"].number_format = "#,##0.00"

dv = DataValidation(type="list", formula1='"Core,Enabling,Optional"', allow_blank=True)
dv.error = "Use Core, Enabling, or Optional"
dv.prompt = "Could-go class"
ws.add_data_validation(dv)
dv.add("G4:G63")

dv2 = DataValidation(
    type="list",
    formula1='"Habitat,Power,Mobility,Life support,Comms,Manufacture,Science,Safety"',
    allow_blank=True,
)
ws.add_data_validation(dv2)
dv2.add("B4:B63")

widths = [8, 16, 32, 16, 16, 14, 16, 16, 24, 36, 14, 16]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.oddHeader.left.text = "Mars 1-lb catalog  ·  n = 1+0"
ws.oddFooter.right.text = "Could-go inventory — not a flight list"

# ============================================================
# Sheet 3 — Category matrix (pivot-like)
# ============================================================
ws2 = wb.create_sheet("Category matrix")
ws2.merge_cells("A1:F1")
ws2["A1"] = "CATEGORY × CLASS MATRIX"
ws2["A1"].font = title_font
ws2["A1"].fill = fill_title
ws2["A1"].alignment = Alignment(horizontal="left", vertical="center")
ws2.row_dimensions[1].height = 26

ws2["A3"] = "Category"
ws2["B3"] = "Core lb"
ws2["C3"] = "Enabling lb"
ws2["D3"] = "Optional lb"
ws2["E3"] = "Total lb"
ws2["F3"] = "Item count"
for c in range(1, 7):
    ws2.cell(3, c).font = header_font
    ws2.cell(3, c).fill = fill_header
    ws2.cell(3, c).alignment = center
    ws2.cell(3, c).border = thin

cats = ["Habitat", "Power", "Mobility", "Life support", "Comms", "Manufacture", "Science", "Safety"]
for i, cat in enumerate(cats):
    r = 4 + i
    ws2.cell(r, 1, cat).font = body_font
    ws2.cell(r, 2, f'=SUMIFS(Catalog!D:D,Catalog!B:B,A{r},Catalog!G:G,"Core")')
    ws2.cell(r, 3, f'=SUMIFS(Catalog!D:D,Catalog!B:B,A{r},Catalog!G:G,"Enabling")')
    ws2.cell(r, 4, f'=SUMIFS(Catalog!D:D,Catalog!B:B,A{r},Catalog!G:G,"Optional")')
    ws2.cell(r, 5, f"=B{r}+C{r}+D{r}")
    ws2.cell(r, 6, f"=COUNTIF(Catalog!B:B,A{r})")
    for c in range(1, 7):
        ws2.cell(r, c).border = thin
        ws2.cell(r, c).font = body_font if c == 1 else black_calc
        if c in (2, 3, 4, 5):
            ws2.cell(r, c).number_format = "#,##0.00"
        ws2.cell(r, c).alignment = left if c == 1 else center
        if i % 2:
            ws2.cell(r, c).fill = fill_alt

ws2["A12"] = "ALL"
ws2["A12"].font = label_font
ws2["B12"] = "=SUM(B4:B11)"
ws2["C12"] = "=SUM(C4:C11)"
ws2["D12"] = "=SUM(D4:D11)"
ws2["E12"] = "=SUM(E4:E11)"
ws2["F12"] = "=SUM(F4:F11)"
for c in range(1, 7):
    ws2.cell(12, c).fill = fill_yellow
    ws2.cell(12, c).border = thin
    ws2.cell(12, c).font = label_font
    if c in (2, 3, 4, 5):
        ws2.cell(12, c).number_format = "#,##0.00"

chart = BarChart()
chart.type = "col"
chart.grouping = "stacked"
chart.title = "Catalog pounds by category and class"
chart.y_axis.title = "Catalog lb"
chart.x_axis.title = None
data = Reference(ws2, min_col=2, min_row=3, max_col=4, max_row=11)
cats_ref = Reference(ws2, min_col=1, min_row=4, max_row=11)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats_ref)
chart.shape = 4
chart.style = 10
chart.width = 18
chart.height = 9
ws2.add_chart(chart, "A15")

for i, w in enumerate([16, 14, 14, 14, 14, 12], start=1):
    ws2.column_dimensions[get_column_letter(i)].width = w

ws2["A32"] = "How to read"
ws2["A32"].font = label_font
ws2.merge_cells("A33:F35")
ws2["A33"] = (
    "Each category is the sum of catalog pounds sitting in that class. "
    "Core = needed for a minimal useful site. Enabling = unlocks Core to actually work. "
    "Optional = quality, science, comfort, extra margin. "
    "None of these rows are a commitment to fly."
)
ws2["A33"].alignment = Alignment(wrap_text=True, vertical="top")
ws2["A33"].font = note_font

# ============================================================
# Sheet 4 — How to add
# ============================================================
ws3 = wb.create_sheet("How to add")
ws3["A1"] = "ADD A ROW WITHOUT BREAKING n = 1 + 0"
ws3["A1"].font = title_font
ws3["A1"].fill = fill_title
ws3.merge_cells("A1:B1")
ws3.row_dimensions[1].height = 26

steps = [
    "1. Next empty n on Catalog is already numbered (51–60 reserved).",
    "2. Fill Category, Item, Catalog lb (blue), Assembly k (blue), class, depends, enables, notes.",
    "3. Do not put anything in row index 0. Do not renumber existing n downward.",
    "4. If an item is retired, keep its n and mark class or notes “retired”. Do not collapse to zero.",
    "5. Depends-on uses other n values (example: 5,6). Use — if it depends on nothing.",
    "6. Category matrix and totals update from the Catalog sheet automatically.",
    "7. Catalog lb may be a decimal (0.25 sensor). Packet ceil always rounds up to a whole pound.",
    "8. This file is the could-go matrix. A later flight pick-list is a filter on class + n, not a new index system.",
]
for i, s in enumerate(steps, start=3):
    ws3.merge_cells(start_row=i, start_column=1, end_row=i, end_column=2)
    ws3.cell(i, 1, s).font = body_font
    ws3.row_dimensions[i].height = 20

ws3.column_dimensions["A"].width = 100
ws3.column_dimensions["B"].width = 20

ws3["A13"] = "Generator check"
ws3["B13"] = "=Algorithm!C9"
ws3["B13"].font = black_calc
ws3["C13"] = "must be 1"
ws3["C13"].font = note_font

from pathlib import Path
out = Path(__file__).resolve().parent / "A-and-P-catalog-matrix.xlsx"
wb.save(out)
print("wrote", out)
